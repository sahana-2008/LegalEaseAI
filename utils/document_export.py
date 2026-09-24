from pathlib import Path
from datetime import datetime


def _output_dir():
    folder = Path("exports")
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def format_docx(content):
    """Create a DOCX document and return its file path."""
    from docx import Document

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = _output_dir() / f"LegalEaseAI_{timestamp}.docx"

    document = Document()
    document.add_heading("LegalEase AI", level=1)

    for paragraph in str(content).split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)

    document.save(file_path)

    return str(file_path)


def format_pdf(content):
    """Create a PDF document and return its file path."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import mm

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = _output_dir() / f"LegalEaseAI_{timestamp}.pdf"

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph("LegalEase AI", styles["Title"]),
        Spacer(1, 12),
    ]

    for paragraph in str(content).split("\n"):
        if paragraph.strip():
            safe_text = (
                paragraph
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            story.append(
                Paragraph(safe_text, styles["BodyText"])
            )
            story.append(Spacer(1, 6))

    document.build(story)

    return str(file_path)


def format_txt(content):
    """Create a TXT document and return its file path."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = _output_dir() / f"LegalEaseAI_{timestamp}.txt"

    with open(file_path, "w", encoding="utf-8") as file:
        file.write(str(content))

    return str(file_path)


def export_as_docx(content, filename="LegalEaseAI_Document.docx"):
    """Export content as DOCX."""
    from docx import Document

    file_path = _output_dir() / filename

    document = Document()
    document.add_heading("LegalEase AI", level=1)

    for paragraph in str(content).split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)

    document.save(file_path)

    return str(file_path)


def export_as_pdf(content, filename="LegalEaseAI_Document.pdf"):
    """Export content as PDF."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    file_path = _output_dir() / filename

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph("LegalEase AI", styles["Title"])
    ]

    for paragraph in str(content).split("\n"):
        if paragraph.strip():
            safe_text = (
                paragraph
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            story.append(
                Paragraph(safe_text, styles["BodyText"])
            )

    document.build(story)

    return str(file_path)


def export_as_txt(content, filename="LegalEaseAI_Document.txt"):
    """Export content as TXT."""

    file_path = _output_dir() / filename

    with open(file_path, "w", encoding="utf-8") as file:
        file.write(str(content))

    return str(file_path)